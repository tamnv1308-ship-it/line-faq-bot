"""Extract Shop details and explicit MST/ST Base cluster memberships, read-only."""
import argparse
import hashlib
import json
import re
from pathlib import Path
from co_regions import REGIONS
from co_store_directory import VERSION,validate,metadata


def code(value):
    if isinstance(value,(float,int)) and not isinstance(value,bool) and value==int(value):
        value=str(int(value))
    text=str(value or '').strip()
    if not re.fullmatch('[0-9]{1,6}',text) or int(text)<=0:
        raise ValueError('Mã shop/base không phải số hợp lệ.')
    return str(int(text))


def text(value):
    return ' '.join(str(value or '').split())


def extract(detail_rows,cluster_rows):
    rows={};memberships={};group_bases={};labels={};base_regions={}
    for row in detail_rows:
        if len(row)<12 or row[2] is None:continue
        key=code(row[2]);region=text(row[0])
        normalized=region[5:] if region.startswith('Vùng ') else region
        if normalized not in REGIONS:continue
        if key in rows:raise ValueError('Mã shop bị trùng trong Chi Tiết: '+key)
        rows[key]=dict(name=text(row[3]),region=region,province=text(row[8]),province_old=text(row[11]),
                       ward=text(row[7]),district_old=text(row[10]),channel=text(row[4]),record_type='detail')
    for row in cluster_rows:
        if len(row)<6 or row[2] is None:continue
        region=text(row[1]);normalized=region[5:] if region.startswith('Vùng ') else region
        if normalized not in REGIONS:continue
        member,base=code(row[2]),code(row[5]);label=text(row[4])
        if not label:raise ValueError('Cụm trống tại shop '+member)
        # A change of capitalization must not invent a second cluster.
        match=re.fullmatch(r'cụm\s+([0-9]+)',label,re.I)
        if match:label='Cụm '+str(int(match.group(1)))
        group=label.casefold()
        if group in group_bases and group_bases[group]!=base:
            raise ValueError('Cụm có nhiều ST base: '+label)
        group_bases[group]=base;labels[group]=label;base_regions.setdefault(base,region)
        # Do not transitively merge clusters through a base that is itself a member.
        for key in (member,base):memberships.setdefault(key,set()).add(group)
        if member not in rows:
            rows[member]=dict(name=text(row[3]),region=region,province='Chưa phân vùng',record_type='base_reference')
    for base,region in base_regions.items():
        rows.setdefault(base,dict(name='',region=region,province='Chưa phân vùng',record_type='base_reference'))
    for key,row in rows.items():
        row['clusters']=[dict(id=g,name=labels[g],base=group_bases[g]) for g in sorted(memberships.get(key,set()))]
    return rows


def workbook_payload(path):
    from openpyxl import load_workbook
    book=load_workbook(path,read_only=True,data_only=True)
    try:
        if not {'Chi Tiết','Cụm'}<=set(book.sheetnames):raise ValueError('Thiếu sheet Chi Tiết/Cụm.')
        details=book['Chi Tiết'];clusters=book['Cụm']
        if details['C2'].value!='Mã chi nhánh' or clusters['C2'].value!='MST' or clusters['E2'].value!='Cụm' or clusters['F2'].value!='ST Base':
            raise ValueError('Tiêu đề/cột workbook đã thay đổi; dừng cập nhật.')
        rows=extract(details.iter_rows(min_row=3,values_only=True),clusters.iter_rows(min_row=3,values_only=True))
    finally:book.close()
    payload=dict(version=VERSION,source=path.name,source_sha256=hashlib.sha256(path.read_bytes()).hexdigest(),stores=rows)
    validate(payload)
    return payload


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--file',required=True,type=Path);parser.add_argument('--output',required=True,type=Path)
    args=parser.parse_args();payload=workbook_payload(args.file)
    args.output.parent.mkdir(parents=True,exist_ok=True)
    args.output.write_text(json.dumps(payload,ensure_ascii=False,separators=(',',':')))
    args.output.chmod(0o600)
    print(json.dumps(metadata(payload),ensure_ascii=False))

if __name__=='__main__':main()
