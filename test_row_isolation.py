import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch

from transfer_parser import parse_form, parse_request
from co_flow import Queue, preview_text
from co_results import items, decode_results, result_text, protected_items, filtered_preview
import mac_worker as w


def item(source='789', product='0131491005415'):
    return dict(source=source,destination='2670',product=product,quantity=1,note='QL cho hàng',status='Mới')


def cells(row, error='', co=''):
    return ['1',row['source']+' - Kho xuất',row['destination']+' - Kho nhận',row['product'],
            'iPhone',str(row['quantity']),'0',row['note'],'',co,error]


def detail(row):
    return dict(product_name='iPhone',source_name=row['source']+' - Kho xuất',
                destination_name=row['destination']+' - Kho nhận')


class ProductLengthTests(unittest.TestCase):
    def form(self,code,source=1340):
        return f'Kho xuất: {source}\nKho nhận: 910\nMSP: {code}\nSL: 2'

    def test_only_exact_thirteen_and_no_padding(self):
        for code in ('131491005413','01314910054133','0003'):
            with self.assertRaisesRegex(ValueError,'đúng 13 số'):parse_form(self.form(code))
        self.assertEqual(parse_form(self.form('0131491005413'))['product'],'0131491005413')

    def test_one_invalid_does_not_stop_valid_rows_or_change_notes(self):
        p=parse_request(self.form('131491005413')+'\n===\n'+self.form('0131491005415',789)+'\nNote: giữ nguyên')
        self.assertEqual(len(items(p)),1)
        self.assertEqual(items(p)[0]['note'],'giữ nguyên')
        self.assertEqual(p['skipped_errors'][0]['product'],'131491005413')
        self.assertIn('12 số',p['skipped_errors'][0]['error'])
        with tempfile.TemporaryDirectory() as tmp:
            from openpyxl import load_workbook
            output=Path(tmp)/'upload.xlsx';w.fill_template(w.TEMPLATE,output,p)
            book=load_workbook(output)
            self.assertEqual(book.active.max_row,2)
            self.assertEqual(book.active['C2'].value,'0131491005415');book.close()

    def test_all_invalid_have_explicit_pair_and_code(self):
        with self.assertRaisesRegex(ValueError,'1340 → 910.*131491005413'):
            parse_request(self.form('131491005413'))


class RowIsolationTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.q=Queue(str(Path(self.tmp.name)/'q.db'))
        self.rows=[item(),item('2320','0131491005413'),item('15877','0131491005417')]

    def test_preview_excludes_failed_row_then_confirms_only_remaining(self):
        p=dict(items=self.rows,status='Mới')
        j=self.q.draft('event','owner','chat',p,True)
        job=self.q.claim(True,batch_protocol=True)
        message=json.dumps(dict(rows=[detail(self.rows[0]),detail(self.rows[2])],
                                excluded=[dict(index=1,error='Mã sản phẩm không tồn tại')]))
        self.assertTrue(self.q.update(j['id'],job['lease'],'preview_ready',message))
        self.assertTrue(self.q.update(j['id'],job['lease'],'preview_ready',message))
        saved=self.q.notifications()[0];payload=json.loads(saved['payload'])
        self.assertEqual([r['source'] for r in items(payload)],['789','15877'])
        self.assertIn('2320 → 2670',preview_text(saved))
        self.q.confirm(j['id'],'owner','chat')
        claimed=self.q.claim(True,j['id'],True)
        self.assertEqual(len(items(claimed['payload'])),2)

    def test_preview_rejects_bad_exclusions_and_mismatched_names(self):
        p=dict(items=self.rows)
        for excluded in ([dict(index=9,error='error')],[dict(index=1,error='')],
                         [dict(index=1,error='error')]*2,[dict(index=True,error='error')]):
            with self.assertRaises(ValueError):filtered_preview(p,json.dumps(dict(rows=[],excluded=excluded)))
        with self.assertRaises(ValueError):
            filtered_preview(p,json.dumps(dict(rows=[detail(self.rows[1]),detail(self.rows[2])],excluded=[dict(index=1,error='error')])))

    def test_only_explicit_failed_rows_release_unknown_inventory(self):
        results=[dict(co='789CO123',error='',state='succeeded'),
                 dict(co='',error='MWG từ chối',state='failed'),
                 dict(co='',error='Cần đối soát',state='unknown')]
        j=self.q.draft('old','owner','chat',dict(items=self.rows))
        with self.q.db() as db:
            db.execute("UPDATE jobs SET state='unknown',result=? WHERE id=?",(json.dumps(results),j['id']))
        new=self.q.draft('new','owner','chat',dict(items=[self.rows[1],dict(self.rows[2],quantity=2),item('999')]),
                         check_duplicates=True)
        payload=json.loads(new['payload'])
        self.assertEqual([r['source'] for r in items(payload)],['2320','999'])
        self.assertIn('đối soát',payload['skipped_errors'][0]['error'])
        self.q.confirm(new['id'],'owner','chat')
        claimed=self.q.claim(batch_protocol=True,recovery_protocol=True,page_ready=True)
        self.assertEqual(claimed['id'],new['id'])
        with self.q.db() as db:old=dict(db.execute('SELECT * FROM jobs WHERE id=?',(j['id'],)).fetchone())
        self.assertEqual(old['state'],'unknown');self.assertEqual(old['result'],json.dumps(results))
        self.assertEqual([r['source'] for r in protected_items(old,True)],['15877'])
        text=result_text(old)
        self.assertIn('✅ 1.',text);self.assertIn('❌ 2.',text);self.assertIn('⚠️ 3.',text)

    def test_legacy_unknown_keeps_every_row_protected(self):
        job=dict(state='unknown',payload=json.dumps(dict(items=self.rows)),result=json.dumps([
            dict(co='',error='uncertain') for _ in self.rows]))
        self.assertEqual(protected_items(job,True),self.rows)
        job['result']='old text';self.assertEqual(protected_items(job,True),self.rows)

    def test_malformed_failure_cannot_release_existing_co(self):
        for result in (dict(co='789CO123',error='err',state='failed'),
                       dict(co='',error='',state='failed'),dict(co='',error='err',state='succeeded')):
            with self.assertRaises(ValueError):decode_results(self.rows[0],json.dumps([result]))


if __name__=='__main__':unittest.main()
