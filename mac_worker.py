"""Mac worker. Browser adapter requires an on-site acceptance test before deployment."""
import argparse
import fcntl
import json
import os
import re
import socket
import time
import threading
import urllib.request
import urllib.error
from pathlib import Path
from co_results import items
from transfer_parser import normalize_status
from worker_runtime import Runtime, atomic_json
from mwg_submit_capture import CreateCapture, merge_results
from co_reconciliation_hold import enforce as enforce_reconciliation_hold, ReconciliationHold
from browser_dialogs import install_dialogs

URL = 'https://inventorytransfers.mwgroup.vn/storechangeordermanually'
ROOT = Path(__file__).resolve().parent / '.co-mac'
TEMPLATE = Path(__file__).resolve().parent / 'templates' / 'StoreChangeOrderTGDD_template.xlsx'
SLOW_LOADING_SECONDS = 15
CREATE_RESULT_WAIT_SECONDS = 30
IDLE_REPAIR_INTERVAL_SECONDS = 15
IDLE_REPAIR_MAX_ATTEMPTS = 5
MAX_PRESERVED_WORKSPACES = 3
RUNTIME = None
TELEMETRY = None

def health_state(state):
    if RUNTIME is not None:RUNTIME.set(state)

def check_stop():
    if RUNTIME is not None:RUNTIME.check_stop()

def browser_wait(page, milliseconds):
    page.wait_for_timeout(milliseconds)
    if RUNTIME is not None:RUNTIME.control_ok()

def dialog_diagnostic(event):
    # The broker supplies only a fixed event name/type and sanitized job ID.
    print('Hộp thoại MWG: '+json.dumps(event,ensure_ascii=False),flush=True)

def progress(job, stage):
    health_state('processing')
    if TELEMETRY is not None:
        try:TELEMETRY.progress(job,stage)
        except Exception:pass
    # Do not log tokens, form notes, or customer data.
    ROOT.mkdir(exist_ok=True)
    line=time.strftime('%Y-%m-%d %H:%M:%S')+' '+job['id']+' '+stage
    print(line,flush=True)
    with (ROOT/'worker-progress.log').open('a') as stream:
        stream.write(line+'\n')

def capture_failure(page, label):
    """Keep local evidence for failures; diagnostics must never block reporting."""
    try:
        if page.is_closed() or page.url.split('?',1)[0].split('#',1)[0].rstrip('/')!=URL:
            return
        folder=ROOT/'diagnostics';folder.mkdir(mode=0o700,exist_ok=True)
        name=re.sub(r'[^a-zA-Z0-9_-]','_',str(label))[:60]
        target=folder/(name+'.png')
        page.screenshot(path=str(target),timeout=2000)
        target.chmod(0o600)
    except Exception:
        pass

def save_create_diagnostic(job, diagnostic):
    """Persist only the allowlisted response shape, never a raw MWG response."""
    folder=ROOT/'diagnostics';folder.mkdir(mode=0o700,exist_ok=True)
    name=re.sub(r'[^a-zA-Z0-9_-]','_',str(job['id']))[:60]
    atomic_json(folder/(name+'-create-response.json'),dict(diagnostic,job_id=job['id'],updated=time.time()))

def upload_excel(page, output):
    # Use the site's import button so its preparation handlers run before upload.
    with page.expect_file_chooser(timeout=15000) as chooser:
        page.get_by_role('button',name=re.compile('Nhập Excel mẫu')).click()
    chooser.value.set_files(str(output))

def import_with_retry(page, job, output, alerts, responses=None, records=None):
    """Use the site's native upload by default; API bridge is explicit opt-in.

    The API preparation callback failed on production before sending the file.
    Keep it out of normal intake until verified against the live widget contract.
    """
    mode=os.environ.get('CO_IMPORT_TRANSPORT','ui').strip().lower()
    if mode not in {'auto','api','ui'}:
        raise RuntimeError('CO_IMPORT_TRANSPORT không hợp lệ; chưa nhập file.')
    if job.get('mode')=='create' and records is not None and responses is not None:
        raise RuntimeError('Nhập file chỉ dùng trước bước xác nhận; chưa tạo CO.')
    if records is not None and responses is not None and mode!='ui':
        from mwg_import_bridge import import_preview, BridgeUnavailable
        if job.get('mode')!='preview':
            raise RuntimeError('API nhập file chỉ dùng trước bước xác nhận; chưa tạo CO.')
        control=page.get_by_role('combobox',name='Vui lòng chọn trạng thái sản phẩm',exact=True)
        label=control.input_value().strip()
        match=re.match(r'^\s*(\d+)\s*-',label)
        if not match or normalize_status(label)!=normalize_status(job['payload'].get('status','Mới')):
            raise RuntimeError('Chưa xác minh được trạng thái MWG; chưa nhập file.')
        alerts.clear();responses.clear()
        try:
            response=import_preview(page,output,records,match.group(1),
                                    on_start=lambda:progress(job,'api_import'))
        except BridgeUnavailable:
            if mode=='api':raise
            progress(job,'api_import_ui_fallback')
        else:
            responses.append(response)
            popup=page.get_by_text('Dữ liệu trong tập tin excel',exact=True)
            deadline=time.monotonic()+7
            while True:
                if alerts:raise RuntimeError('MWG báo lỗi nhập Excel: '+'; '.join(alerts)[-1800:])
                if popup.is_visible():return popup
                if time.monotonic()>=deadline:
                    raise RuntimeError('API đã nhập file nhưng bảng xem trước MWG chưa hiện; chưa tạo CO.')
                browser_wait(page,100)
    # UI retry remains only for sites without the compatible API bridge.

    for attempt in range(2):
        if attempt:
            progress(job,'retry_import_after_timeout')
            page.reload(wait_until='domcontentloaded',timeout=30000)
            prepare_import_controls(page,job)
        alerts.clear()
        requests=set()
        def on_request(request):
            from mwg_import_api import IMPORT_URL
            if (request.method=='POST' and
                    request.url.split('?',1)[0].rstrip('/').lower()==IMPORT_URL.lower()):
                requests.add(request)
        def on_response(response):
            if response.request in requests:responses.append(response)
        if responses is not None:
            responses.clear()
            page.on('request',on_request)
            page.on('response',on_response)
        try:
            progress(job,'upload_excel')
            upload_excel(page,output)
            popup=page.get_by_text('Dữ liệu trong tập tin excel',exact=True)
            deadline=time.monotonic()+15
            while True:
                if alerts:
                    raise RuntimeError('MWG báo lỗi khi nhập Excel: '+'; '.join(alerts)[-1800:])
                if popup.is_visible():
                    return popup
                if time.monotonic()>=deadline:
                    break
                browser_wait(page,250)
        finally:
            if responses is not None:
                for event,handler in (('request',on_request),('response',on_response)):
                    try:page.remove_listener(event,handler)
                    except Exception:pass
    raise RuntimeError('MWG chưa trả kết quả nhập Excel sau 2 lượt chờ, mỗi lượt 15 giây; chưa bấm Tạo CO. Yêu cầu đã dừng.')

def received_import_errors(responses, records):
    from mwg_import_api import import_row_errors
    if len(responses)!=1:
        raise RuntimeError('Chưa nhận được duy nhất phản hồi nhập file từ MWG; chưa tạo CO.')
    response=responses[0]
    if response.status!=200:
        raise RuntimeError('MWG trả lỗi nhập file HTTP '+str(response.status)+'; chưa tạo CO.')
    try:body=response.json()
    except Exception:
        raise RuntimeError('MWG không trả dữ liệu nhập file hợp lệ; kiểm tra phiên đăng nhập. Chưa tạo CO.') from None
    try:return import_row_errors(body,records)
    except (ValueError,TypeError,KeyError):
        # Log schema only, never the response contents or authentication headers.
        schema={'body_type':type(body).__name__}
        if isinstance(body,dict):
            schema['keys']=list(body)
            if isinstance(body.get('Data'),dict):schema['data_keys']=list(body['Data'])
        print('Phản hồi nhập MWG chưa khớp: '+json.dumps(schema,ensure_ascii=False),flush=True)
        raise RuntimeError('Phản hồi nhập MWG không khớp đủ mã kho, mã sản phẩm và số lượng; chưa tạo CO.') from None

def concurrent_rejection(alerts, data):
    for message in alerts:
        normalized=' '.join(message.split())
        pair=r'\[\s*'+re.escape(data['source'])+r'\s*-\s*'+re.escape(data['product'])+r'\s*\]'
        if ('Đang có Kho + MSP đang tạo YCCK' in normalized and
                'không được phép cùng lúc tránh vượt tồn' in normalized and
                re.search(pair,normalized)):
            return message[:2500]
    return None






def result_state(rows):
    if any(row.get('state')=='unknown' or (row['co'] and row['error']) for row in rows):return 'unknown'
    return 'succeeded' if any(row['co'] for row in rows) else 'failed'

def api(route, data, timeout=30):
    base = os.environ['CO_BOT_URL'].rstrip('/')
    if not base.startswith('https://'):
        raise RuntimeError('CO_BOT_URL must use HTTPS')
    req = urllib.request.Request(base+'/co/worker/'+route, data=None if data is None else json.dumps(data).encode(),
        headers={'Content-Type':'application/json', 'Authorization':'Bearer '+os.environ['CO_WORKER_TOKEN']})
    with urllib.request.urlopen(req, timeout=timeout) as response:
        return json.load(response)

def claim_with_reconnect(data):
    # Retry polling only; never replay browser actions or CO submission.
    failures=0
    while True:
        check_stop()
        try:
            request=dict(data,health=RUNTIME.snapshot()) if RUNTIME is not None else data
            result=api('claim',request)
            if failures:
                print('Đã kết nối lại máy chủ; tiếp tục trên phiên Chrome hiện tại.',flush=True)
            return result
        except urllib.error.HTTPError as error:
            if error.code not in (408,429,500,502,503,504):
                raise
        # On Python 3.9 socket.timeout is not a subclass of TimeoutError.
        # HTTPS response reads can raise it directly, without a URLError wrapper.
        except (urllib.error.URLError, socket.timeout, TimeoutError, ConnectionError):
            pass
        failures+=1
        health_state('reconnecting')
        delay=min(5*failures,30)
        print('Tạm mất kết nối máy chủ. Giữ nguyên Chrome, thử lại sau '+str(delay)+' giây.',flush=True)
        time.sleep(delay)

class ResultDeliveryBlocked(RuntimeError):
    pass

def reconcile_result(record):
    snapshot=api('result-status',{'id':record['id'],'lease':record['lease']})
    if not isinstance(snapshot,dict) or not snapshot.get('found'):
        raise ResultDeliveryBlocked('Không tìm thấy yêu cầu trên máy chủ. Giữ kết quả trên Mac để đối soát.')
    if (snapshot.get('lease_matches') is True and snapshot.get('state')==record['state']
            and snapshot.get('result')==record.get('result','')):
        return
    # Only a known pre-submit failure may be retired after cancellation/expiry.
    # Success, unknown outcomes, and submitting barriers must never be discarded.
    if (record['state']=='failed' and record.get('before_submit') is True
            and snapshot.get('state') in {'failed','cancelled','expired'}
            and not re.search(r'[0-9A-Z]+CO[0-9]+',str(snapshot.get('result',''))+record.get('result',''))):
        folder=ROOT/'result-conflicts';folder.mkdir(mode=0o700,exist_ok=True)
        destination=folder/(str(time.time_ns())+'.json')
        with destination.open('x') as stream:
            os.chmod(destination,0o600)
            json.dump({'record':record,'server':snapshot},stream,ensure_ascii=False)
            stream.flush();os.fsync(stream.fileno())
        print('Đã đối chiếu '+record['id']+': lỗi trước khi tạo CO; máy chủ đã kết thúc yêu cầu. Đã lưu bằng chứng và gỡ kẹt.',flush=True)
        return
    raise ResultDeliveryBlocked('Kết quả '+record['id']+' không khớp máy chủ ('+str(snapshot.get('state'))+'). Giữ kết quả trên Mac; cần đối soát, không tạo lại CO.')

def deliver_result(record):
    """Resend the same durable result only; never repeat MWG creation."""
    failures=0
    while True:
        try:
            try:
                reply=api('result',record)
            except urllib.error.HTTPError as error:
                if error.code==409:
                    reconcile_result(record)
                    return
                if error.code not in (408,429,500,502,503,504):
                    raise ResultDeliveryBlocked('Máy chủ từ chối kết quả (HTTP '+str(error.code)+'). Giữ kết quả trên Mac; cần kiểm tra kết nối/cấu hình.') from error
                raise
            if TELEMETRY is not None and isinstance(reply,dict):
                try:TELEMETRY.acknowledge(reply.get('telemetry_ack'))
                except Exception:pass
            if not isinstance(reply,dict) or reply.get('ok') is not True:
                raise RuntimeError('Máy chủ chưa xác nhận đã lưu kết quả')
            if failures:
                print('Đã gửi lại kết quả thành công. Giữ nguyên mã CO và phiên Chrome.',flush=True)
            return
        except ResultDeliveryBlocked as error:
            if TELEMETRY is not None:
                try:TELEMETRY.delivery_error(record,error.__cause__ or error)
                except Exception:pass
            health_state('blocked')
            raise
        except Exception as error:
            if TELEMETRY is not None:
                try:TELEMETRY.delivery_error(record,error)
                except Exception:pass
            health_state('reconnecting')
            check_stop()  # Durable result remains on disk; never replay MWG submission.
            failures+=1
            delay=min(5*failures,30)
            detail='HTTP '+str(error.code) if isinstance(error,urllib.error.HTTPError) else type(error).__name__
            print('Chưa gửi được kết quả ('+detail+'). Đã lưu trên Mac; thử lại sau '+str(delay)+' giây. Không tạo lại CO.',flush=True)
            time.sleep(delay)

class RecoveryGuard:
    """Allow independent work only after this page is visibly and network idle."""
    def __init__(self,page,track_submissions=True):
        self.page=page
        self.track_submissions=track_submissions
        self.pending=set()
        self.last_reason=None
        self.loading_since=None
        self.repair_failed=False
        self.repair_at=None
        self.repair_attempts=0
        self.retirement_ready=False
        self.retirement_check_at=0
        # Historical unknown jobs stay protected by the server's stock checks.
        # Only an uncertain submission on THIS page must prevent page repair.
        self.uncertain_job=None
        self.late_submissions=[]
        page.on('request',self.request)
        page.on('requestfinished',self.finished)
        page.on('requestfailed',self.finished)
    def request(self,request):
        if request.resource_type in ('xhr','fetch'):
            self.pending.add(request)
            self.retirement_ready=False
    def finished(self,request):
        self.pending.discard(request)
    def ready(self):
        reason=self.readiness_reason()
        if reason is None:
            self.loading_since=None
            self.repair_failed=False
            self.repair_at=None
            self.repair_attempts=0
        if reason!=self.last_reason:
            if reason:print('MWG chưa sẵn sàng: '+reason,flush=True)
            self.last_reason=reason
        return reason is None
    def page_health(self, ready, prepared_id=None):
        """Describe loading without treating a deferred repair as a failure."""
        if ready:
            self.loading_since=None
            self.repair_failed=False
            self.repair_at=None
            self.repair_attempts=0
            return 'awaiting_confirmation' if prepared_id else 'ready'
        if self.page.is_closed():return 'blocked'
        if self.page.url.split('?',1)[0].split('#',1)[0].rstrip('/')!=URL:
            self.loading_since=None
            return 'login'
        if self.uncertain_job is not None or self.late_submissions or self.repair_failed:
            return 'blocked'
        now=time.monotonic()
        if self.loading_since is None:self.loading_since=now
        return 'loading_slow' if now-self.loading_since>=SLOW_LOADING_SECONDS else 'loading'
    def readiness_reason(self,ignore_submission=False):
        page=self.page
        if not ignore_submission and (self.uncertain_job is not None or self.late_submissions):
            return 'Đang chờ đối soát phản hồi tạo CO; giữ nguyên trang.'
        try:
            if page.is_closed():return 'Cửa sổ đã đóng.'
            if page.url.split('?',1)[0].split('#',1)[0].rstrip('/')!=URL:
                return 'Chưa ở trang tạo CO; kiểm tra đăng nhập MWG.'
            if self.pending:return 'Còn kết nối dữ liệu đang chờ MWG trả lời.'
            if page.locator('.k-loading-mask:visible, .blockUI:visible, [aria-busy="true"]:visible').count():
                return 'Trang đang có lớp chờ tải dữ liệu.'
            if any(control.is_visible() for control in page.get_by_text(re.compile(r'^\s*Loading(?:\.\.\.)?\s*$',re.I)).all()):
                return 'Trang đang hiển thị Loading.'
            button=page.get_by_role('button',name=re.compile('Nhập Excel mẫu'))
            status=page.get_by_role('combobox',name='Vui lòng chọn trạng thái sản phẩm',exact=True)
            if not button.is_visible() or not button.is_enabled():return 'Nút Nhập Excel mẫu chưa sẵn sàng.'
            if not status.is_visible():return 'Chưa hiển thị ô trạng thái sản phẩm.'
            if normalize_status(status.input_value(timeout=1000)) is None:return 'Ô trạng thái sản phẩm trống hoặc chưa nhận diện được.'
            if page.get_by_text('Dữ liệu trong tập tin excel',exact=True).is_visible():return 'Bảng nhập Excel vẫn đang mở.'
            return None
        except Exception as error:
            return 'Không đọc được trạng thái trang ('+type(error).__name__+').'
    def wait_ready(self, timeout=15):
        deadline=time.monotonic()+timeout
        stable=None
        while time.monotonic()<deadline:
            if self.ready():
                if stable is None: stable=time.monotonic()
                elif time.monotonic()-stable>=1: return True
            else:
                stable=None
            try:
                browser_wait(self.page,250)
            except Exception as error:
                if self.page.is_closed() or 'has been closed' in str(error):
                    return False
                raise
        return False

def report(job, state, result='', *, before_submit=False):
    if TELEMETRY is not None:
        try:
            TELEMETRY.finish(job,result if state in {'failed','unknown'} else '')
            TELEMETRY.progress(job,'send_result')
        except Exception:pass
    record = {'id':job['id'], 'lease':job['lease'], 'state':state, 'result':result}
    if TELEMETRY is not None:
        try:record['telemetry']=TELEMETRY.batch(job)
        except Exception:pass
    if before_submit:
        record['before_submit']=True
    pending = ROOT/'pending.json'
    tmp = ROOT/'pending.tmp'
    with tmp.open('w') as f:
        json.dump(record, f); f.flush(); os.fsync(f.fileno())
    tmp.replace(pending)
    deliver_result(record)
    pending.unlink()
    if TELEMETRY is not None:
        try:TELEMETRY.finish(job)
        except Exception:pass
    progress(job,'reported_'+state)


def archive_settled_submission(guard, job, capture, previous):
    """Release a browser resource, never the server's unresolved stock hold.

    No in-flight request may be retired. Persist evidence and confirm that the
    server still holds this job before allowing a completed page to be closed.
    Keep capture listeners until close, so a delayed second POST revokes this.
    """
    if guard.retirement_ready:return
    if time.monotonic()<guard.retirement_check_at:return
    guard.retirement_check_at=time.monotonic()+15
    if (ROOT/'pending.json').exists():return
    closed=capture.page.is_closed()
    if not closed and (not capture.done or guard.pending):return
    if not closed and guard.readiness_reason(ignore_submission=True) is not None:return
    try:
        snapshot=api('result-status',{'id':job['id'],'lease':job['lease']},timeout=5)
        if not isinstance(snapshot,dict) or not snapshot.get('found'):return
        if snapshot.get('state') not in {'unknown','succeeded','failed','cancelled','expired'}:return
        # Owner cancellation clears the lease but retains the exact unknown result.
        cancelled=(snapshot.get('state')=='cancelled' and
                   json.loads(snapshot.get('result') or 'null')==previous)
        if snapshot.get('lease_matches') is not True and not cancelled:return
        rows=merge_results(previous,capture.rows) if capture.rows is not None else previous
        folder=ROOT/'settled-submissions';folder.mkdir(mode=0o700,exist_ok=True)
        evidence={'version':1,'id':job['id'],'archived_at':time.time(),
                  'request_count':len(capture.requests),'page_closed':closed,
                  'transport_finished':bool(capture.done),'problem':capture.problem,
                  'response_diagnostic':capture.diagnostic,
                  'previous_outcomes':previous,'observed_outcomes':rows,
                  'server_state':snapshot['state'],'server_result':snapshot.get('result',''),
                  'records':[{k:row.get(k) for k in ('source','destination','product','quantity')}
                             for row in items(job['payload'])],
                  'replay_allowed':False}
        atomic_json(folder/(job['id']+'.json'),evidence)
        capture_failure(capture.page,job['id']+'-settled')
        # Browser events can be delivered during the screenshot/status request.
        if (len(capture.requests)!=evidence['request_count'] or
                (not capture.page.is_closed() and
                 (guard.pending or not capture.done or guard.readiness_reason(ignore_submission=True) is not None))):return
        guard.retirement_ready=True
        progress(job,'settled_submission_archived')
    except Exception as error:
        print('Chưa lưu/đối chiếu được trang CO '+job['id']+' ('+type(error).__name__+'); giữ nguyên trang.',flush=True)


def reconcile_late_submissions(guard):
    for entry in list(guard.late_submissions):
        job,capture,previous=entry
        if not capture.done and not capture.page.is_closed():continue
        if capture.rows is not None and len(capture.requests)==1:
            resolved=merge_results(previous,capture.rows)
            if all(row.get('state') in {'succeeded','failed'} for row in resolved):
                progress(job,'late_create_response_verified')
                report(job,result_state(resolved),json.dumps(resolved,ensure_ascii=False))
                if guard.uncertain_job==job['id']:guard.uncertain_job=None
                capture.close();guard.late_submissions.remove(entry)
                continue
        archive_settled_submission(guard,job,capture,previous)


def workspace_can_retire(guard):
    if (ROOT/'pending.json').exists():return False
    closed=guard.page.is_closed()
    if guard.pending and not closed:return False
    if not guard.uncertain_job and not guard.late_submissions:return True
    if not closed and guard.readiness_reason(ignore_submission=True) is not None:return False
    return (guard.retirement_ready and
            all(capture.done or capture.page.is_closed() for _,capture,_ in guard.late_submissions))


def retire_workspace(guard):
    if not workspace_can_retire(guard):return False
    try:
        if not guard.page.is_closed():guard.page.close(run_before_unload=False)
    except Exception:return False
    for _,capture,_ in guard.late_submissions:capture.close()
    guard.late_submissions.clear()
    for event,handler in (('request',guard.request),('requestfinished',guard.finished),('requestfailed',guard.finished)):
        try:guard.page.remove_listener(event,handler)
        except Exception:pass
    return True


def prune_preserved_workspaces(preserved, prepared_id=None):
    if prepared_id or (ROOT/'pending.json').exists():return
    for saved in list(preserved):
        if retire_workspace(saved):preserved.remove(saved)

def isolate_workspace(context, page, guard, preserved, prepared_id=None):
    if prepared_id or (ROOT/'pending.json').exists():
        return page,guard,False
    prune_preserved_workspaces(preserved)
    can_retire=workspace_can_retire(guard)
    if not can_retire and len(preserved)>=MAX_PRESERVED_WORKSPACES:
        return page,guard,False
    # Create the replacement before closing the last working browser page.
    fresh=context.new_page()
    retired=retire_workspace(guard) if can_retire else False
    if not retired and len(preserved)>=MAX_PRESERVED_WORKSPACES:
        fresh.close(run_before_unload=False)
        return page,guard,False
    if not retired:preserved.append(guard)
    return fresh,RecoveryGuard(fresh,track_submissions=guard.track_submissions),True

def fill_template(source, target, data):
    from openpyxl import load_workbook
    wb = load_workbook(source)
    ws = wb.active
    if [ws.cell(1,c).value for c in range(1,6)] != ['Kho xuất','Kho nhận hàng','Mã sản phẩm','Số lượng','Ghi chú']:
        raise RuntimeError('Cấu trúc mẫu Excel đã thay đổi.')
    if any(c.value is not None for row in ws.iter_rows(min_row=2) for c in row):
        raise RuntimeError('Mẫu có dữ liệu sẵn; dừng để tránh nhập nhầm.')
    records=items(data)
    last_row=len(records)+1
    if ws.max_row>last_row:
        ws.delete_rows(last_row+1,ws.max_row-last_row)
    for row_number in list(ws.row_dimensions):
        if row_number>last_row:
            del ws.row_dimensions[row_number]
    for row_number,record in enumerate(records,2):
        for col,key in enumerate(['source','destination','product','quantity','note'],1):
            cell=ws.cell(row_number,col); cell.value=record[key]
            if isinstance(record[key],str):
                cell.data_type='s'
                cell.number_format='@'
    wb.save(target)

def table_with(page, header):
    return page.get_by_role('table').filter(has=page.get_by_role('columnheader',name=header,exact=True))


def data_rows(table):
    result=[]
    for row in table.get_by_role('row').all():
        # Main CO table: read visible cells, not hidden Angular template columns.
        # The import popup is validated separately from its HTTP response.
        cells=row.get_by_role('cell').all_text_contents()
        if cells and any(value.strip() for value in cells):
            result.append((row,cells))
    return result

def match_rows(rows, records):
    if len(rows)!=len(records):
        return None
    matched=[]; used=set()
    for data in records:
        candidates=[]
        for index,(_,cells) in enumerate(rows):
            if index in used or len(cells)!=11:
                continue
            if (cells[3].strip()==data['product'] and cells[5].strip()==str(data['quantity'])
                    and cells[7].strip()==data['note']
                    and all(re.match(r'^'+re.escape(value)+r'(?:\s*-|\s*$)',actual.strip())
                            for value,actual in [(data['source'],cells[1]),(data['destination'],cells[2])])):
                candidates.append((index,cells))
        if len(candidates)!=1:
            return None
        index,cells=candidates[0]; used.add(index); matched.append(cells)
    return matched

def wait_matched_rows(page, table, records, timeout=7):
    """After accepting import, MWG can briefly show old, partial or unnamed rows."""
    deadline=time.monotonic()+timeout
    stable=None
    signature=None
    while True:
        rows=data_rows(table)
        matched=match_rows(rows,records)
        now=time.monotonic()
        if matched is not None and all(cells[4].strip() for cells in matched):
            current=tuple(tuple(cells) for cells in matched)
            if current!=signature:
                signature=current;stable=now
            elif now-stable>=0.3:
                return matched
        else:
            stable=None;signature=None
        if now>=deadline:
            # Only counts and field comparisons, never notes, names or tokens.
            diagnostic={'expected_rows':len(records),'observed_rows':len(rows),
                        'cell_counts':[len(cells) for _,cells in rows],
                        'matches':[]}
            for data in records:
                candidates=[cells for _,cells in rows if len(cells)==11 and cells[3].strip()==data['product']]
                diagnostic['matches'].append([{
                    'source':bool(re.match(r'^'+re.escape(data['source'])+r'(?:\s*-|\s*$)',cells[1].strip())),
                    'destination':bool(re.match(r'^'+re.escape(data['destination'])+r'(?:\s*-|\s*$)',cells[2].strip())),
                    'quantity':cells[5].strip()==str(data['quantity']),
                    'note':cells[7].strip()==data['note'],
                    'name':bool(cells[4].strip())} for cells in candidates])
            print('Đối chiếu bảng CO chưa khớp: '+json.dumps(diagnostic,ensure_ascii=False),flush=True)
            raise RuntimeError('Dữ liệu trên trang chưa khớp đủ '+str(len(records))+' dòng sau '+str(timeout)+' giây; chưa tạo CO.')
        browser_wait(page,100)

def cancel_import(page):
    cancel=page.get_by_text('Hủy',exact=True)
    if cancel.is_visible():
        cancel.click()

def wait_ui(page, ready, message, timeout=15):
    deadline=time.monotonic()+timeout
    while not ready():
        if time.monotonic()>=deadline:
            raise RuntimeError(message)
        browser_wait(page,100)


def select_brands(page):
    desired={'1 - thegioididong','2 - dienmayxanh','16 - TopZone'}
    control=page.locator('input[role="listbox"]')
    control.wait_for(state='visible',timeout=15000)
    owns=(control.get_attribute('aria-owns') or '').split()
    if owns and {value.strip() for value in page.locator('[id="'+owns[-1]+'"] .k-state-selected').all_text_contents()}==desired:
        return  # Reuse verified selections without opening/reloading the menu.
    control.click()
    wait_ui(page,lambda: bool(control.get_attribute('aria-owns')),
            'Danh sách thương hiệu chưa sẵn sàng; chưa nhập file.')
    owns=(control.get_attribute('aria-owns') or '').split()
    if not owns:
        raise RuntimeError('Danh sách thương hiệu chưa sẵn sàng; chưa nhập file.')
    choices=page.locator('[id="'+owns[-1]+'"]')
    for label in desired:
        choices.get_by_role('option',name=label,exact=True).wait_for(state='attached',timeout=15000)
    for option in choices.get_by_role('option').all():
        label=option.inner_text().strip()
        selected='k-state-selected' in (option.get_attribute('class') or '').split()
        if selected != (label in desired):
            if not option.is_visible():
                control.click()
                option.wait_for(state='visible',timeout=15000)
            option.click()
            wait_ui(page,lambda: ('k-state-selected' in (option.get_attribute('class') or '').split()) == (label in desired),
                    'Thương hiệu chưa cập nhật; chưa nhập file.')
    selected={text.strip() for text in choices.locator('.k-state-selected').all_text_contents()}
    if selected != desired:
        raise RuntimeError('Không chọn đúng ba thương hiệu TGDD, DMX, TopZone.')
    control.press('Escape')

def select_product_status(page, status):
    requested=normalize_status(status)
    if requested is None:
        raise RuntimeError('Trạng thái yêu cầu không hợp lệ; chưa nhập file.')
    control=page.get_by_role('combobox',name='Vui lòng chọn trạng thái sản phẩm',exact=True)
    control.wait_for(state='visible',timeout=15000)
    if normalize_status(control.input_value())==requested:
        return
    control.press('Alt+ArrowDown')
    matches=[option for option in page.get_by_role('option').all()
             if option.is_visible() and normalize_status(option.inner_text())==requested]
    if len(matches)!=1:
        raise RuntimeError('Không tìm thấy duy nhất trạng thái '+status+' trên MWG; chưa nhập file.')
    matches[0].click()
    wait_ui(page,lambda: normalize_status(control.input_value())==requested,
            'Không xác nhận được trạng thái '+status+'; chưa nhập file.')
    if normalize_status(control.input_value())!=requested:
        raise RuntimeError('Không xác nhận được trạng thái '+status+'; chưa nhập file.')

def wait_product_status(page, control, timeout=7):
    # Require a nonempty value to stay stable for one second, not a transient render.
    deadline=time.monotonic()+timeout
    previous=None
    stable_since=None
    while time.monotonic()<deadline:
        value=control.input_value(timeout=max(1,int((deadline-time.monotonic())*1000))).strip() if control.is_visible() else ''
        now=time.monotonic()
        if normalize_status(value) is not None and value==previous:
            if stable_since is not None and now-stable_since>=1:
                return True
        else:
            previous=value
            stable_since=now if value else None
        browser_wait(page,500)
    return False


def prepare_import_controls(page, job):
    # MWG loads inventory statuses from the selected brands. Waiting for a
    # status before selecting brands can deadlock a freshly loaded page.
    control=page.get_by_role('combobox',name='Vui lòng chọn trạng thái sản phẩm',exact=True)
    brand_timeout=False
    for attempt in range(5):
        progress(job,'select_brands')
        try:
            select_brands(page)
            brand_timeout=False
        except Exception as error:
            if type(error).__name__!='TimeoutError':
                raise
            brand_timeout=True
            if attempt<4:
                progress(job,'reload_brand_timeout')
                page.reload(wait_until='domcontentloaded',timeout=30000)
                continue
            break
        if wait_product_status(page,control,timeout=7):
            progress(job,'select_product_status')
            select_product_status(page,job['payload'].get('status','Mới'))
            return
        if attempt<4:
            progress(job,'reload_blank_product_status')
            page.reload(wait_until='domcontentloaded',timeout=30000)
    if brand_timeout:
        raise RuntimeError('Danh sách thương hiệu MWG chưa tải đầy đủ sau 5 lượt kiểm tra; chưa nhập file, chưa tạo CO. Vui lòng thử lại khi MWG tải đầy đủ.')
    raise RuntimeError('Ô trạng thái sản phẩm vẫn chưa sẵn sàng sau 5 lượt tải (chờ trạng thái tối đa 7 giây mỗi lượt); chưa nhập file, chưa tạo CO. Vui lòng thử lại khi MWG tải đầy đủ.')

def process(page, job, prepared_id=None, guard=None):
    submitting=False
    creation=None
    alerts=[]
    records=items(job['payload'])
    try:
        enforce_reconciliation_hold(ROOT,job['id'],records)
    except ReconciliationHold as error:
        progress(job,'reconciliation_hold')
        report(job,'failed',str(error),before_submit=True)
        return
    original_records=records
    outcomes=[{'co':'','error':''} for _ in records]
    excluded=[]
    alerts=install_dialogs(page.context,dialog_diagnostic).begin(page,job['id'])
    try:
        reuse=job.get('mode')=='create' and prepared_id==job['id']
        if job.get('mode')=='create' and job['payload'].get('product_name') and not reuse:
            raise RuntimeError('Trang chờ xác nhận không còn trong phiên bot này; chưa tạo CO. Gửi lại yêu cầu để nhập file mới.')
        if reuse:
            progress(job,'reuse_imported_row')
            if page.url.split('?',1)[0].split('#',1)[0].rstrip('/')!=URL:
                raise RuntimeError('Trang MWG đã thay đổi khi chờ xác nhận; chưa tạo CO.')
        else:
            progress(job,'reuse_page_before_request')
            if page.url.split('?',1)[0].split('#',1)[0].rstrip('/')!=URL:
                page.goto(URL,wait_until='domcontentloaded',timeout=30000)
            page.get_by_role('button',name='Tạo CO',exact=True).wait_for(timeout=30000)
            prepare_import_controls(page,job)
            progress(job,'prepare_local_template')
            if not TEMPLATE.is_file():
                raise RuntimeError('Thiếu mẫu Excel lưu trên Mac; cần cập nhật templates/StoreChangeOrderTGDD_template.xlsx. Chưa tạo CO.')
            folder=ROOT/job['id']; folder.mkdir(exist_ok=True)
            output=folder/'StoreChangeOrderTGDD_upload.xlsx'
            fill_template(TEMPLATE,output,job['payload'])
            original_indices=list(range(len(records)))
            while True:
                responses=[]
                popup=import_with_retry(page,job,output,alerts,responses=responses,records=records)
                errors=received_import_errors(responses,records)
                progress(job,'import_response_verified')
                progress(job,'accept_import')
                if not any(errors):break
                cancel_import(page)
                keep=[];indices=[]
                for index,data,error in zip(original_indices,records,errors):
                    if error:excluded.append({'index':index,'error':error[:500]})
                    else:keep.append(data);indices.append(index)
                if not keep:
                    failures={row['index']:row['error'] for row in excluded}
                    report(job,'failed',json.dumps([{'co':'','error':failures[n],'state':'failed'}
                        for n in range(len(original_records))],ensure_ascii=False),before_submit=True)
                    return
                # This only repeats the import, before confirmation or any CO click.
                records=keep;original_indices=indices
                progress(job,'skip_invalid_import_rows')
                fill_template(TEMPLATE,output,dict(job['payload'],items=records))
            page.get_by_text('Đồng ý',exact=True).click()
            popup.wait_for(state='hidden',timeout=15000)
        table=table_with(page,'Mã yêu cầu chuyển kho')
        matched=wait_matched_rows(page,table,records)
        for cells in matched:
            if cells[9].strip() or cells[10].strip():
                raise RuntimeError('Trang đã có mã CO hoặc lỗi; dừng để kiểm tra.')
        names=[cells[4].strip() for cells in matched]
        if not all(names):
            raise RuntimeError('MWG chưa trả đủ tên sản phẩm; chưa tạo CO.')
        if job.get('mode')=='preview':
            preview={'rows':[{'product_name':cells[4].strip(),'source_name':cells[1].strip(),'destination_name':cells[2].strip()} for cells in matched]}
            if excluded:preview['excluded']=excluded
            report(job,'preview_ready',json.dumps(preview,ensure_ascii=False))
            return job['id']
        if excluded:
            raise RuntimeError('Dòng nhập đã thay đổi; cần gửi bản xác nhận mới trước khi tạo CO.')
        for data,name in zip(records,names):
            if data.get('product_name') and name!=data['product_name']:
                raise RuntimeError('Tên sản phẩm trên MWG khác bản đã xác nhận; chưa tạo CO.')
        data=job['payload']
        if reuse:
            actual=page.get_by_role('combobox',name='Vui lòng chọn trạng thái sản phẩm',exact=True).input_value().strip()
            if normalize_status(actual)!=normalize_status(data.get('status','Mới')):
                raise RuntimeError('Trạng thái sản phẩm đã thay đổi khi chờ xác nhận; chưa tạo CO.')
            control=page.locator('input[role="listbox"]')
            owns=(control.get_attribute('aria-owns') or '').split()
            selected=set() if not owns else {x.strip() for x in page.locator('[id="'+owns[-1]+'"] .k-state-selected').all_text_contents()}
            if selected!={'1 - thegioididong','2 - dienmayxanh','16 - TopZone'}:
                raise RuntimeError('Thương hiệu đã thay đổi khi chờ xác nhận; chưa tạo CO.')
        if guard is not None:guard.uncertain_job=job['id']
        report(job,'submitting')
        submitting=True
        progress(job,'submit_co')
        alerts.clear()
        creation=CreateCapture(page,records,on_diagnostic=lambda evidence:save_create_diagnostic(job,evidence))
        create_button=page.get_by_role('button',name='Tạo CO',exact=True)
        creation.arm_click(create_button)
        click_error=None
        try:create_button.click()
        except Exception as error:click_error=error
        # Start at the trusted DOM click (or observed POST), excluding preparation,
        # server reporting and actionability waits before the actual click.
        deadline=creation.result_deadline(CREATE_RESULT_WAIT_SECONDS)
        if click_error is not None and creation.clicked_at is None and not creation.requests:
            raise click_error
        while True:
            if creation.rows is not None and len(creation.requests)==1:
                outcomes=merge_results(outcomes,creation.rows)
                if all(row.get('state') in {'succeeded','failed'} for row in outcomes):
                    progress(job,'create_response_verified')
                    state=result_state(outcomes)
                    report(job,state,json.dumps(outcomes,ensure_ascii=False))
                    if guard is not None:guard.uncertain_job=None
                    return
            current=match_rows(data_rows(table),records)
            if current is not None:
                displayed=[dict(row) for row in outcomes]
                for index,cells in enumerate(current):
                    co=cells[9].strip(); error=cells[10].strip()
                    if co or error:
                        if co and not re.fullmatch(r'[0-9A-Z]+CO[0-9]+',co):
                            continue
                        previous_co=outcomes[index]['co']
                        if previous_co and co!=previous_co:
                            co=previous_co
                            error='Kết quả MWG thay đổi sau khi đã nhận mã CO; cần đối soát. '+error
                        displayed[index]={'co':co,'error':error[:350],
                            'state':'unknown' if co and error else 'succeeded' if co else 'failed'}
                for index,data in enumerate(records):
                    rejection=concurrent_rejection(alerts,data)
                    if rejection and not outcomes[index]['co'] and not outcomes[index]['error']:
                        displayed[index]={'co':'','error':rejection[:500],'state':'failed'}
                outcomes=merge_results(outcomes,displayed)
                if all(row.get('state') in {'succeeded','failed'} for row in outcomes) and len(creation.requests)<=1:
                    state=result_state(outcomes)
                    report(job,state,json.dumps(outcomes,ensure_ascii=False))
                    if guard is not None and state!='unknown':guard.uncertain_job=None
                    return
            remaining=deadline-time.monotonic()
            if remaining<=0:break
            browser_wait(page,max(1,min(300,int(remaining*1000))))
        raise RuntimeError(creation.problem or f'MWG chưa trả mã CO hoặc lỗi sau {CREATE_RESULT_WAIT_SECONDS} giây; tiếp tục theo dõi phản hồi, không tạo lại.')
    except Exception as error:
        progress(job,'exception_'+type(error).__name__)
        capture_failure(page,job['id'])
        if (ROOT/'pending.json').exists():
            raise # Preserve undelivered outcome, including successful CO codes.
        message=str(error).splitlines()[0][:1800]
        if alerts and not (submitting and any(concurrent_rejection(alerts,data) for data in records)):
            message+=' | '+'; '.join(alerts)[-500:]
        if 'phiên làm việc đã hết hạn' in message.lower():
            message=('Phiên đăng nhập MWG trên Mac đã hết hạn. '
                     'Đăng nhập lại MWG trong cửa sổ Chrome của bot. '
                     'Bot tự nhận lại khi trang sẵn sàng. '
                     'Sau đó gửi lại form và xác nhận. Không cần đổi thông tin kho hoặc sản phẩm.')
        elif 'has been closed' in message:
            message=('Chrome của bot đã đóng hoặc mất kết nối. Giữ cửa sổ Chrome của bot mở. '
                     + ('Cần kiểm tra MWG trước khi gửi lại vì đã bắt đầu bước tạo CO.' if submitting else
                        'Chưa tạo CO; hãy gửi lại form và xác nhận yêu cầu mới.'))
        elif 'strict mode violation' in message:
            message='Bot gặp lỗi chọn thành phần trên trang MWG; chưa hoàn tất thao tác. Cần kiểm tra bot trước khi gửi lại.'
        if not submitting:
            try:
                cancel_import(page)
            except Exception:
                pass
        if submitting:
            if creation is not None:
                stage=('create_no_request_observed' if not creation.requests else
                       'create_response_unresolved' if creation.done else 'create_response_pending')
                progress(job,stage)
            for outcome in outcomes:
                if not outcome['co'] and not outcome['error']:
                    outcome['error']='Chưa rõ kết quả; cần đối soát MWG. '+message[:250]
                    outcome['state']='unknown'
            report(job,'unknown',json.dumps(outcomes,ensure_ascii=False))
            if guard is not None and creation is not None:
                guard.late_submissions.append((job,creation,outcomes))
                creation=None  # Keep observing this exact request after the result wait expires.
        else:
            if excluded:
                failures={row['index']:row['error'] for row in excluded}
                message=json.dumps([{'co':'','error':failures.get(index,message),'state':'failed'}
                                    for index in range(len(original_records))],ensure_ascii=False)
            report(job,'failed',message,before_submit=True)
    finally:
        if creation is not None:creation.close()

def open_browser(pw, background=False):
    # The bundled CfT archive lacks valid macOS signing resources on this Mac.
    # Use the installed signed Chrome and preserve the original bot login profile.
    if not Path('/Applications/Google Chrome.app/Contents/MacOS/Google Chrome').is_file():
        raise RuntimeError('Không tìm thấy Google Chrome trong Applications. Chưa nhận yêu cầu CO.')
    context=pw.chromium.launch_persistent_context(
        str(ROOT/'browser'),channel='chrome',headless=background,accept_downloads=True,
        args=['--disable-backgrounding-occluded-windows','--disable-renderer-backgrounding'])
    install_dialogs(context,dialog_diagnostic)
    if RUNTIME is not None:
        runtime=RUNTIME
        generation=runtime.browser_opened()
        context.on('close',lambda *_: runtime.browser_closed(generation))
    # MWG may request focus after navigation; leave foreground choice to the user.
    context.add_init_script("window.focus = () => {};")
    return context, context.pages[0] if context.pages else context.new_page()

def refresh_idle(page, guard=None):
    # Preserve the current result page when reconciliation is still required.
    if (ROOT/'pending.json').exists():
        return False
    if guard is not None and (guard.uncertain_job is not None or guard.late_submissions or guard.pending or guard.repair_failed):
        return False
    states=api('status',None)['states']
    if any(states.get(state,0) for state in ('draft','queued','verified_queued','preview_queued','running','submitting','preview_running')):
        return False
    if (guard is None or not guard.track_submissions) and states.get('unknown',0):
        return False
    if (ROOT/'pending.json').exists() or (guard is not None and
            (guard.pending or guard.uncertain_job is not None or guard.late_submissions or guard.repair_failed)):
        return False
    page.reload(wait_until='domcontentloaded',timeout=30000)
    print('Đã tải lại trang MWG sau 5 phút khi rảnh.',flush=True)
    if guard is not None:
        # Reload clears the brands that populate the dependent status control.
        # Prepare them in this idle cycle, rather than leaving intake disabled
        # until a later repair tick and potentially reloading a loading status.
        deadline=time.monotonic()+7
        while guard.pending and time.monotonic()<deadline:
            browser_wait(page,100)
        if (guard.pending or guard.uncertain_job is not None or guard.late_submissions
                or (ROOT/'pending.json').exists() or page.is_closed()
                or page.url.split('?',1)[0].split('#',1)[0].rstrip('/')!=URL):
            return True
        states=api('status',None,timeout=5)['states']
        if any(states.get(state,0) for state in ('draft','queued','verified_queued','preview_queued','running','submitting','preview_running')):
            return True
        if guard.pending or guard.uncertain_job is not None or guard.late_submissions or (ROOT/'pending.json').exists():
            return True
        try:
            select_brands(page)
            guard.wait_ready(timeout=7)
        except Exception as error:
            print('MWG chưa hoàn tất chuẩn bị sau tải lại ('+type(error).__name__+'); giữ trang chờ phục hồi.',flush=True)
    return True

def ensure_browser(pw, context, page, background=False):
    if not page.is_closed():
        return context, page
    # A closed tab can leave the context alive. A closed window cannot.
    try:
        return context, context.new_page()
    except Exception as error:
        if 'has been closed' not in str(error):
            raise
        return open_browser(pw, background=background)

def find_idle_mwg_page(context, page, guard, prepared_id=None, excluded_pages=()):
    """Follow a login-created tab only when no current CO state can be lost."""
    if prepared_id or guard.uncertain_job is not None or guard.late_submissions:
        return page
    if (ROOT/'pending.json').exists() or guard.pending:
        return page
    def on_mwg(candidate):
        return (not candidate.is_closed() and
                candidate.url.split('?',1)[0].split('#',1)[0].rstrip('/')==URL)
    try:
        if on_mwg(page):return page
        candidates=[candidate for candidate in context.pages if candidate is not page
                    and not any(candidate is excluded for excluded in excluded_pages) and on_mwg(candidate)]
        # Multiple open CO forms are ambiguous; never pick a form arbitrarily.
        if len(candidates)!=1:return page
        states=api('status',None)['states']
        if any(states.get(state,0) for state in ('draft','preview_running','running','submitting')):
            return page
        if not guard.track_submissions and states.get('unknown',0):return page
        if (ROOT/'pending.json').exists():return page
        print('Đã nhận đúng thẻ tạo CO sau đăng nhập; giữ nguyên phiên MWG.',flush=True)
        return candidates[0]
    except Exception:
        # Discovery is optional: a closed tab or unavailable status cannot stop the worker.
        return page

def start_heartbeat():
    stop=threading.Event()
    def pulse():
        previous=None
        sent=0
        while not stop.is_set():
            RUNTIME.persist()
            health=RUNTIME.snapshot()
            if health['state']!=previous or time.monotonic()-sent>=5:
                if health['state']!=previous and health['state'] in {'driver_error','control_stalled'}:
                    print('Phát hiện lỗi điều khiển: '+health['state']+'. Giữ nguyên trang và yêu cầu, không tự tạo lại CO.',flush=True)
                try:
                    payload={'online':True,'health':health}
                    if TELEMETRY is not None:
                        try:payload['telemetry']=TELEMETRY.batch()
                        except Exception:pass
                    answer=api('heartbeat',payload,timeout=5)
                    if TELEMETRY is not None and isinstance(answer,dict):
                        try:TELEMETRY.acknowledge(answer.get('telemetry_ack'))
                        except Exception:pass
                    previous=health['state'];sent=time.monotonic()
                except Exception:
                    print('Mất kết nối heartbeat; máy chủ tự ngừng nhận CO khi quá 30 giây.',flush=True)
            stop.wait(1)
    thread=threading.Thread(target=pulse,daemon=True)
    thread.start()
    def shutdown():
        stop.set()
        thread.join(timeout=6)
        try:api('heartbeat',{'online':False,'health':{'state':'stopped','progress_age_seconds':0}},timeout=5)
        except Exception:pass
    return shutdown


def recover_idle_controls(page, guard, prepared_id=None):
    """One idle recovery step, spaced 15s apart and capped per loading episode.

    The main loop keeps pumping browser events between steps. A pending request
    consumes a readiness check, never permission to reload a busy CO page.
    """
    def protected():
        return (prepared_id or guard.uncertain_job is not None or guard.late_submissions
                or (ROOT/'pending.json').exists() or page.is_closed()
                or page.url.split('?',1)[0].split('#',1)[0].rstrip('/')!=URL)
    def server_idle():
        states=api('status',None,timeout=5)['states']
        return (not any(states.get(k,0) for k in ('draft','verified_queued','preview_running','running','submitting'))
                and (guard.track_submissions or not states.get('unknown',0)))
    if protected():return False
    if guard.ready():return True
    if guard.repair_failed:return False
    now=time.monotonic()
    if guard.repair_at is None:
        guard.repair_at=now+IDLE_REPAIR_INTERVAL_SECONDS
        return False
    if now<guard.repair_at:return False
    # Also back off failed status reads; never spin or mutate when idle is unknown.
    guard.repair_at=now+IDLE_REPAIR_INTERVAL_SECONDS
    try:
        check_stop()
        if not server_idle() or protected():return False
        if guard.repair_attempts>=IDLE_REPAIR_MAX_ATTEMPTS:
            # The fifth attempt receives the same settling time before giving up.
            guard.repair_failed=True
            print('MWG chưa sẵn sàng sau 5 lượt thử cách nhau 15 giây. Tạm ngừng nhận CO; giữ Chrome để kiểm tra.',flush=True)
            capture_failure(page,'readiness-error')
            return False
        guard.repair_attempts+=1
        attempt=str(guard.repair_attempts)+'/5'
        if guard.pending:
            print('Kiểm tra MWG lượt '+attempt+': còn phản hồi dữ liệu, giữ nguyên trang.',flush=True)
            return False
        print('Phục hồi MWG khi rảnh lượt '+attempt+'.',flush=True)
        try:
            # Brands must be selected before the dependent status list can load.
            select_brands(page)
        except Exception as error:
            print('Chưa chuẩn bị được thương hiệu MWG ('+type(error).__name__+').',flush=True)
        if guard.ready():return True
        # Selecting brands starts asynchronous GetInventoryStatus requests. Give
        # them a bounded settling budget; an immediate reload cancels that work.
        if not protected() and guard.wait_ready(timeout=7):return True
        # Recheck after UI waits: confirmation/submission may have started meanwhile.
        if protected() or guard.pending or not server_idle() or protected() or guard.pending:
            return False
        page.reload(wait_until='domcontentloaded',timeout=15000)
        return guard.ready()
    finally:
        if guard.repair_at is not None:
            guard.repair_at=time.monotonic()+IDLE_REPAIR_INTERVAL_SECONDS

def main():
    global RUNTIME, TELEMETRY
    env_file=Path(__file__).resolve().parent/'.env'
    if env_file.exists():
        for line in env_file.read_text().splitlines():
            key,sep,value=line.partition('=')
            if sep and key in {'CO_BOT_URL','CO_WORKER_TOKEN'}:
                os.environ.setdefault(key,value.strip())
    parser=argparse.ArgumentParser()
    parser.add_argument('--login',action='store_true')
    parser.add_argument('--live',action='store_true')
    parser.add_argument('--background',action='store_true',help='Chạy Chrome không hiện cửa sổ sau khi đăng nhập')
    parser.add_argument('--api',action='store_true',help='Dùng API MWG sau xác nhận')
    parser.add_argument('--probe-import',action='store_true',help='Thử API nhập file, không tạo CO')
    args=parser.parse_args()
    if args.probe_import and args.live:
        parser.error('Không chạy thử API đồng thời với nhận lệnh thật.')
    if not (args.login or args.live or args.probe_import):
        parser.error('Dùng --login để đăng nhập; --live để xử lý tác vụ đã xác nhận.')
    os.umask(0o077); ROOT.mkdir(exist_ok=True)
    lock=open(ROOT/'worker.lock','a')
    try:
        fcntl.flock(lock,fcntl.LOCK_EX|fcntl.LOCK_NB)
    except BlockingIOError:
        print('BOT ĐANG MỞ ở phiên khác. Không mở thêm. Dùng Kiem-tra-BOT.command để xem trạng thái.',flush=True)
        raise SystemExit(3)
    os.environ['PLAYWRIGHT_BROWSERS_PATH']=str(Path(__file__).resolve().parent/'.pw-browsers')
    from playwright.sync_api import sync_playwright
    from worker_observability import Telemetry
    try:TELEMETRY=Telemetry(ROOT) if args.live else None
    except Exception:TELEMETRY=None
    RUNTIME=Runtime(ROOT);RUNTIME.install()
    shutdown=start_heartbeat() if args.live else lambda:None
    try:
        with sync_playwright() as pw:
            RUNTIME.watch_driver()
            context,page=open_browser(pw, background=args.background and not args.login)
            try:
                if args.login:
                    page.goto(URL)
                    RUNTIME.wait_login()
                    if args.background:
                        # Close cleanly to flush the persistent login before reopening it.
                        # This happens before claiming any request or importing a file.
                        context.close()
                        context,page=open_browser(pw, background=True)
                if args.probe_import:
                    from mwg_import_api import probe_import
                    if page.url.split('?',1)[0].rstrip('/') != URL:
                        page.goto(URL,wait_until='domcontentloaded',timeout=30000)
                    page.get_by_role('button',name='Tạo CO',exact=True).wait_for(timeout=30000)
                    print('Đã mở MWG. Đang tìm mã xác thực...',flush=True)
                    token_control=page.locator('input[name="__RequestVerificationToken"]')
                    if token_control.count()!=1:
                        raise ValueError('Không tìm thấy duy nhất một mã xác thực trên trang MWG.')
                    token=token_control.first.input_value(timeout=5000)
                    sample={'source':'852','destination':'10606','product':'0131491005414',
                            'quantity':2,'note':'KIEM TRA API NHAP FILE - KHONG TAO CO','status':'Mới'}
                    output=ROOT/'api-probe.xlsx'
                    fill_template(TEMPLATE,output,sample)
                    started=time.monotonic()
                    print('Đang gửi API nhập Excel (tối đa 60 giây). Không tạo CO.',flush=True)
                    rows=probe_import(context.request,token,output,[sample],status_id='1')
                    elapsed=round(time.monotonic()-started,3)
                    result={'ok':True,'operation':'import_only','created_co':False,
                            'seconds':elapsed,'rows':[{'source':r['FROMSTOREID'],
                            'destination':r['TOSTOREID'],'product':r['PRODUCTID'],
                            'name':r['PRODUCTNAME'],'quantity':r['QUANTITY']} for r in rows]}
                    (ROOT/'api-probe-result.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
                    print('API NHẬP FILE THÀNH CÔNG: '+str(elapsed)+' giây. Đã đối chiếu đúng dữ liệu. KHÔNG tạo CO.',flush=True)
                if args.live:
                    print('Bot đang chạy. Dùng Kiem-tra-BOT.command để xem trạng thái; Tat-BOT.command để dừng.', flush=True)
                    pending=ROOT/'pending.json'
                    if pending.exists():
                        record=json.loads(pending.read_text()); deliver_result(record); pending.unlink()
                        if record['state']=='submitting':
                            raise RuntimeError('Có tác vụ dở ở bước tạo. Kiểm tra MWG trước khi chạy lại.')
                    refresh_at=time.monotonic()+300
                    prepared_id=None
                    preserved_guards=[]
                    from mwg_api_worker import ApiWorker
                    api_worker=ApiWorker() if args.api else None
                    guard=RecoveryGuard(page,track_submissions=not args.api)
                    recovery_ready=False
                    while True:
                        check_stop()
                        for preserved in preserved_guards:
                            reconcile_late_submissions(preserved)
                        reconcile_late_submissions(guard)
                        prune_preserved_workspaces(preserved_guards,prepared_id)
                        if guard.uncertain_job and not prepared_id and not pending.exists():
                            # Preserve the exact submission page and its response listeners.
                            # Server unknown-stock protection prevents replay on the fresh page.
                            page,guard,isolated=isolate_workspace(context,page,guard,preserved_guards,prepared_id)
                            if isolated:
                                recovery_ready=False
                                print('Đã tách yêu cầu cần đối soát; dùng trang riêng cho yêu cầu khác, không tạo lại CO cũ.',flush=True)
                        context,page=ensure_browser(pw,context,page,background=args.background)
                        page=find_idle_mwg_page(context,page,guard,prepared_id,
                                                excluded_pages=[saved.page for saved in preserved_guards])
                        if page.url in ('about:blank','chrome://newtab/'):
                            try:
                                health_state('loading')
                                page.goto(URL,wait_until='domcontentloaded',timeout=15000)
                            except Exception:
                                print('MWG chưa mở được. Giữ Chrome và thử lại; chưa lấy yêu cầu.',flush=True)
                                time.sleep(5)
                                continue
                        if guard.page is not page:
                            guard=RecoveryGuard(page,track_submissions=not args.api)
                            recovery_ready=False
                        ready=guard.ready()
                        health_state(guard.page_health(ready,prepared_id))
                        if not ready:
                            on_mwg=page.url.split('?',1)[0].split('#',1)[0].rstrip('/')==URL
                            if on_mwg and not prepared_id:
                                try:
                                    ready=recover_idle_controls(page,guard,prepared_id)
                                    if ready:refresh_at=time.monotonic()+300
                                except Exception:
                                    ready=False
                            health_state(guard.page_health(ready,prepared_id))
                        job=claim_with_reconnect({'preview_protocol':1,'prepared_id':prepared_id,'batch_protocol':1,
                                                 'reset_protocol':1,'recovery_protocol':int(not args.api),
                                                 'page_ready':recovery_ready and guard.ready()})
                        if job and job.get('reset_workspace'):
                            page,guard,isolated=isolate_workspace(context,page,guard,preserved_guards,prepared_id)
                            if isolated:
                                recovery_ready=False
                                refresh_at=time.monotonic()+300
                                print('Đã nhận reset từ chủ bot; giữ bằng chứng và trang CO còn phản hồi.',flush=True)
                            else:
                                print('Chưa reset: cần đối soát các trang đang giữ hoặc kết quả chưa gửi.',flush=True)
                            continue
                        if job and job.get('recovery_required'):
                            recovery_ready=guard.wait_ready()
                            if not recovery_ready:
                                print('MWG còn bận hoặc chưa sẵn sàng. Giữ trang để đối soát; chưa lấy yêu cầu khác.',flush=True)
                                time.sleep(5)
                            continue
                        if job and job.get('release_preview'):
                            prepared_id=None
                            if api_worker: api_worker.prepared=None
                            refresh_at=time.monotonic()+300
                            continue
                        if job:
                            health_state('processing')
                            recovery_ready=False
                            if api_worker:
                                import sys
                                prepared_id=api_worker.process(page,job,prepared_id,sys.modules[__name__])
                            else:
                                prepared_id=process(page,job,prepared_id,guard=guard)
                            refresh_at=time.monotonic()+300
                        elif ready and not prepared_id and time.monotonic()>=refresh_at:
                            try:
                                health_state('loading')
                                refresh_idle(page,guard=guard)
                            except Exception as error:
                                print('Không tải lại được trang MWG: '+type(error).__name__,flush=True)
                            refresh_at=time.monotonic()+300
                        if not job:
                            # Pump Playwright events so completion of XHR can clear the guard.
                            try:browser_wait(page,500 if prepared_id else 1000)
                            except Exception:pass
            finally:
                # Flush the persistent profile even on a cooperative stop or an error.
                try:context.close()
                except Exception:pass
    finally:
        shutdown()
        RUNTIME.close()
        print('BOT ĐÃ TẮT. Chrome của bot và kết nối nhận CO đã dừng.',flush=True)
        RUNTIME=None;TELEMETRY=None
        lock.close()

if __name__=='__main__':
    main()
