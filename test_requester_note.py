import json
import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook, Workbook
from co_flow import add_requester_note, Queue, PRIVATE_CO_OWNER
from mac_worker import fill_template
from unittest.mock import Mock

class RequesterNoteTests(unittest.TestCase):
    def item(self,note=''):
        return dict(source='1',destination='2',product='003',quantity=1,note=note,status='Mới')
    def test_empty_existing_and_repeat(self):
        for original in ('','abc\nGiữ nguyên'):
            p=self.item(original)
            expected=original+(' ' if original else '')+'Mọi thắc mắc liên hệ Tâm'
            add_requester_note(p,'Tâm');add_requester_note(p,'Tâm')
            self.assertEqual(p['note'],expected)
    def test_limit_and_batch_atomicity(self):
        suffix='Mọi thắc mắc liên hệ Tâm'
        p=self.item('x'*(499-len(suffix)))
        add_requester_note(p,'Tâm');self.assertEqual(len(p['note']),500)
        p=dict(items=[self.item('abc'),self.item('x'*500)])
        with self.assertRaisesRegex(ValueError,'500'):add_requester_note(p,'Tâm')
        self.assertEqual(p['items'][0]['note'],'abc')
    def test_missing_name_does_not_expose_id(self):
        for name in (None,'','Không lấy được tên','U'+'a'*32):
            with self.assertRaisesRegex(ValueError,'tên hiển thị LINE'):add_requester_note(self.item(),name)
    def test_admin_confirmation_and_retry_preserve_requester(self):
        with tempfile.TemporaryDirectory() as folder:
            q=Queue(str(Path(folder)/'q.db'))
            p=add_requester_note(self.item('abc'),'Người gửi')
            first=q.draft('event','sender','chat',p)
            again=q.draft('event','sender','chat',p)
            self.assertEqual(first['id'],again['id'])
            q.confirm(first['id'],PRIVATE_CO_OWNER,'chat')
            with q.db() as db:job=dict(db.execute('SELECT * FROM jobs WHERE id=?',(first['id'],)).fetchone())
            self.assertEqual(job['owner'],'sender')
            self.assertEqual(json.loads(job['payload'])['note'],'abc Mọi thắc mắc liên hệ Người gửi')
    def test_shared_excel_note(self):
        with tempfile.TemporaryDirectory() as folder:
            p=add_requester_note(self.item('abc'),'Tâm')
            target=Path(folder)/'upload.xlsx'
            template=Path(folder)/'template.xlsx'
            wb=Workbook();wb.active.append(['Kho xuất','Kho nhận hàng','Mã sản phẩm','Số lượng','Ghi chú']);wb.save(template)
            fill_template(template,target,p)
            self.assertEqual(load_workbook(target).active.cell(2,5).value,p['note'])

    def test_mac_api_payload(self):
        try:
            from mwg_api_worker import create
        except ImportError:
            self.skipTest('API adapter is shipped separately on the Mac worker')
        p=add_requester_note(self.item('abc'),'Tâm')
        row=dict(FROMSTOREID=1,TOSTOREID=2,INVENTORYSTATUSID=1,PRODUCTID='003',QUANTITY=1,NOTE=p['note'],STORECHANGEORDERTYPEID=0,CONTENT='',ISURGENT=0)
        request=Mock();request.post.return_value.status=200
        create(request,'dummy-token',[row])
        self.assertEqual(json.loads(request.post.call_args.kwargs['data']['Json'])[0]['NOTE'],p['note'])
if __name__=='__main__':unittest.main()
