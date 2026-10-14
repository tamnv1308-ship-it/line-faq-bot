"""A fixed intake window groups messages without changing confirmation authority."""
import json
import os
import tempfile
import time
import types
import unittest
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest.mock import Mock, patch

from co_flow import Queue, DuplicateRequest, PRIVATE_CO_OWNER, install
from co_results import items, result_text, line_text_messages


def payload(number=1, status='Mới'):
    return dict(source=str(100+number), destination='999999',
                product='0131491005418', quantity=1, status=status,
                note='QL cho hàng '+str(number), requester_name='Người thử nghiệm')


class GatherTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.q=Queue(str(Path(self.tmp.name)/'queue.db'))
        self.now=Mock(return_value=1000)
        self.clock=patch('co_flow.time',types.SimpleNamespace(time=self.now,monotonic=time.monotonic))
        self.clock.start()
        self.addCleanup(self.clock.stop)

    def draft(self, number, **kwargs):
        options=dict(check_product=True, check_duplicates=True, gather_seconds=10,
                     reply_token='token'+str(number))
        options.update(kwargs)
        return self.q.draft('event'+str(number),options.pop('owner','owner'),
                            options.pop('chat','chat'),options.pop('payload',payload(number)),**options)

    def test_fixed_window_and_one_file_with_original_rows(self):
        first=self.draft(1)
        self.now.return_value=1009.99
        second=self.draft(2)
        self.assertEqual(first['id'],second['id'])
        self.assertIsNone(self.q.claim(True,batch_protocol=True))
        self.now.return_value=1010
        job=self.q.claim(True,batch_protocol=True)
        self.assertEqual(job['mode'],'preview')
        self.assertEqual([r['source'] for r in items(job['payload'])],['101','102'])
        self.assertEqual([r['note'] for r in items(job['payload'])],['QL cho hàng 1','QL cho hàng 2'])
        from mac_worker import fill_template,TEMPLATE
        from openpyxl import load_workbook
        file=Path(self.tmp.name)/'batch.xlsx'
        fill_template(TEMPLATE,file,job['payload'])
        workbook=load_workbook(file)
        self.addCleanup(workbook.close)
        self.assertEqual(workbook.active.max_row,3)
        self.assertEqual(workbook.active['C3'].value,'0131491005418')
        self.assertEqual(workbook.active['E3'].value,'QL cho hàng 2')

    def test_exact_deadline_starts_new_batch_even_if_mac_is_busy(self):
        first=self.draft(1)
        self.now.return_value=1010
        second=self.draft(2)
        self.assertNotEqual(first['id'],second['id'])

    def test_owner_chat_and_status_are_separate(self):
        jobs=[self.draft(1),self.draft(2,owner='another'),
              self.draft(3,chat='another'),self.draft(4,payload=payload(4,'Đã sử dụng'))]
        self.assertEqual(len({j['id'] for j in jobs}),4)
        merged=self.draft(5,payload=payload(5,'Đã sử dụng'))
        self.assertEqual(merged['id'],jobs[-1]['id'])

    def test_max_seven_rows_without_splitting_incoming_message(self):
        first=self.draft(1,payload={'items':[payload(n) for n in range(1,7)],'status':'Mới'})
        self.now.return_value=1001
        another=self.draft(7,payload={'items':[payload(7),payload(8)],'status':'Mới'})
        self.assertNotEqual(first['id'],another['id'])
        filled=self.draft(9)
        self.assertEqual(filled['id'],first['id'])
        self.assertEqual(len(items(json.loads(filled['payload']))),7)

    def test_redelivery_retains_batch_and_latest_reply_across_restart(self):
        first=self.draft(1)
        self.now.return_value=1002
        second=self.draft(2)
        self.q=Queue(self.q.path)
        self.now.return_value=1003
        self.assertEqual(self.draft(2)['id'],first['id'])
        self.assertEqual(self.draft(1)['id'],first['id'])
        with self.q.db() as db:
            stored=db.execute('SELECT payload FROM jobs WHERE id=?',(first['id'],)).fetchone()
        self.assertEqual(len(items(json.loads(stored['payload']))),2)
        self.assertEqual(json.loads(stored['payload'])['request_count'],2)
        self.assertEqual(self.q.take_preview_reply(first['id'])['token'],'token2')

    def test_all_duplicates_skipped_without_changing_original_batch(self):
        first=self.draft(1)
        self.assertEqual(self.draft(2,payload=payload(1))['state'],'skipped')
        with self.q.db() as db:
            p=json.loads(db.execute('SELECT payload FROM jobs WHERE id=?',(first['id'],)).fetchone()['payload'])
        self.assertEqual(len(items(p)),1)
        self.assertEqual(self.q.take_preview_reply(first['id'])['token'],'token1')

    def test_partial_duplicates_append_only_new_rows_and_retain_warning(self):
        first=self.draft(1)
        merged=self.draft(2,payload={'items':[payload(1),payload(2)]})
        self.assertEqual(merged['id'],first['id'])
        p=json.loads(merged['payload'])
        self.assertEqual([r['source'] for r in items(p)],['101','102'])
        self.assertEqual(len(p['skipped_duplicates']),1)
        merged=self.draft(3)
        self.assertEqual(len(json.loads(merged['payload'])['skipped_duplicates']),1)

    def test_cancel_blocks_entire_batch_and_does_not_absorb_future_messages(self):
        first=self.draft(1);self.draft(2)
        self.assertIn('Không tìm thấy',self.q.confirm(first['id'],'another','chat',cancel=True))
        self.q.confirm(first['id'],'owner','chat',cancel=True)
        self.assertEqual(self.draft(2)['state'],'cancelled')
        self.assertNotEqual(self.draft(3)['id'],first['id'])

    def test_preview_single_confirmation_and_partial_results(self):
        first=self.draft(1);self.draft(2)
        self.now.return_value=1010
        self.assertIsNone(self.q.claim(True)) # Older worker cannot receive multiple rows.
        job=self.q.claim(True,batch_protocol=True)
        self.assertTrue(self.q.update(job['id'],job['lease'],'preview_ready',json.dumps(['Phone 1','Phone 2'])))
        confirmation=self.q.notifications()[0]
        self.assertIn('Phone 2',self.q.confirmation(confirmation))
        self.assertEqual(len(line_text_messages(self.q.confirmation(confirmation))[-1]['quickReply']['items']),2)
        self.q.mark_notified(job['id'],'draft')
        self.assertIsNone(self.q.claim(True,job['id'],True))
        self.assertIn('Không tìm thấy',self.q.confirm(job['id'],'owner','other-chat'))
        self.q.confirm(job['id'],PRIVATE_CO_OWNER,'owner-private-chat')
        create=self.q.claim(True,job['id'],True)
        self.assertEqual(create['mode'],'create')
        self.assertEqual(len(items(create['payload'])),2)
        self.q.update(create['id'],create['lease'],'submitting')
        self.q.update(create['id'],create['lease'],'succeeded',json.dumps([
            {'co':'101CO123','error':''},{'co':'','error':'Không đủ tồn'}]))
        result=self.q.notifications()[0]
        self.assertIn('1/2 có mã CO',result_text(result))
        self.assertIn('Không đủ tồn',result_text(result))
        self.assertEqual(self.draft(2)['state'],'succeeded')

    def test_confirmation_ttl_starts_after_notification_and_expires_whole_batch(self):
        first=self.draft(1);self.draft(2)
        self.now.return_value=1010
        job=self.q.claim(True,batch_protocol=True)
        self.q.update(job['id'],job['lease'],'preview_ready','["Phone 1","Phone 2"]')
        self.now.return_value=1020
        self.q.mark_notified(job['id'],'draft')
        self.now.return_value=1079
        self.assertIsNone(self.q.claim(True,job['id'],True))
        self.now.return_value=1080
        self.assertEqual(self.q.claim(True,job['id'],True),{'release_preview':True})
        self.assertEqual(self.draft(2)['state'],'expired')

    def test_concurrent_intake_and_claim_are_serialized(self):
        with ThreadPoolExecutor(max_workers=4) as pool:
            jobs=list(pool.map(self.draft,range(1,8)))
        self.assertEqual(len({j['id'] for j in jobs}),1)
        self.now.return_value=1010
        with ThreadPoolExecutor(max_workers=4) as pool:
            claims=list(pool.map(lambda _:self.q.claim(True,batch_protocol=True),range(4)))
        claims=[c for c in claims if c]
        self.assertEqual(len(claims),1)
        self.assertEqual(len(items(claims[0]['payload'])),7)

    def test_claimed_preview_never_changes(self):
        first=self.draft(1)
        self.now.return_value=1010
        job=self.q.claim(True,batch_protocol=True)
        later=self.draft(2)
        self.assertNotEqual(first['id'],later['id'])
        self.assertEqual(len(items(job['payload'])),1)

    def test_cleanup_counts_messages_and_tombstones_every_event(self):
        from co_summary import aggregate,cleanup
        from co_summary_view import details
        first=self.draft(1);second=self.draft(2)
        self.q.confirm(first['id'],'owner','chat',cancel=True)
        with self.q.db() as db:job=dict(db.execute('SELECT * FROM jobs WHERE id=?',(first['id'],)).fetchone())
        self.assertEqual(aggregate([job],{})[0]['requests'],2)
        from datetime import timezone
        self.assertEqual(details([job],{},timezone.utc)['requests']['owner'],2)
        # Cleanup runs against calendar time; these test jobs date back to 1970.
        cleanup(self.q,Path(self.tmp.name)/'images')
        with self.q.db() as db:
            self.assertEqual(db.execute('SELECT COUNT(*) AS n FROM co_gather_events').fetchone()['n'],0)
            self.assertEqual(db.execute('SELECT COUNT(*) AS n FROM co_gather_windows').fetchone()['n'],0)
        for n in (1,2):
            with self.assertRaises(DuplicateRequest):self.draft(n)


class GatherWebhookTests(unittest.TestCase):
    def test_authenticated_status_exposes_effective_gather_setting(self):
        with tempfile.TemporaryDirectory() as tmp:
            routes={}
            def route(path):
                def register(func):
                    routes[path]=func
                    return func
                return register
            app=types.SimpleNamespace(post=route,get=route)
            flask=types.SimpleNamespace(request=types.SimpleNamespace(headers={'Authorization':'Bearer '+'x'*32}),
                                        jsonify=lambda **kwargs:kwargs,abort=lambda code:None)
            env={'CO_ENABLED':'1','CO_DB_PATH':str(Path(tmp)/'q.db'),'CO_DATABASE_URL':'','CO_WORKER_TOKEN':'x'*32}
            with patch.dict(os.environ,env),patch.dict('sys.modules',{'flask':flask}):
                for setting in ('0','10'):
                    os.environ['CO_GATHER_SECONDS']=setting
                    install(app,lambda *a:None,lambda *a:None)
                    status=routes['/co/worker/status']()
                    self.assertEqual(status['gather_seconds'],int(setting))
                    self.assertEqual(status['max_batch_rows'],7)
                    self.assertEqual(status['copied_text_version'],1)

    def test_default_ten_seconds_from_line_to_one_confirmation(self):
        with tempfile.TemporaryDirectory() as tmp:
            path=str(Path(tmp)/'q.db');replies=[]
            env={'CO_ENABLED':'1','CO_DB_PATH':path,'CO_DATABASE_URL':'','CO_WORKER_TOKEN':'x'*32,
                 'CO_ALLOWED_USER_IDS':'owner','CO_ACK_RECEIPT':'0'}
            app=types.SimpleNamespace(post=lambda _:lambda f:f,get=lambda _:lambda f:f,
                                      logger=types.SimpleNamespace(warning=lambda *args:None))
            flask=types.SimpleNamespace(request=None,jsonify=None,abort=None)
            with patch.dict(os.environ,env),patch.dict('sys.modules',{'flask':flask}),patch('co_flow.time.time',return_value=2000) as now:
                os.environ.pop('CO_GATHER_SECONDS',None)
                handle,notify=install(app,lambda token,text:replies.append((token,text)) or True,
                                      lambda *args:False,requester_name=lambda e:'Người thử nghiệm')
                q=Queue(path);q.heartbeat()
                for n in (1,2):
                    event=types.SimpleNamespace(reply_token='reply'+str(n),webhook_event_id='e'+str(n),
                        source=types.SimpleNamespace(user_id='owner',group_id='chat'),
                        message=types.SimpleNamespace(id=str(n),text=f'Kho xuất: {100+n}\nKho nhận: 999999\nMSP: 0131491005418'))
                    self.assertTrue(handle(event))
                self.assertIsNone(q.claim(True,batch_protocol=True))
                now.return_value=2010
                work=q.claim(True,batch_protocol=True)
                self.assertEqual(len(items(work['payload'])),2)
                q.update(work['id'],work['lease'],'preview_ready','["Phone 1","Phone 2"]')
                notify()
                self.assertEqual(len(replies),1)
                self.assertEqual(replies[0][0],'reply2')
                self.assertIn('Phone 2',replies[0][1])
                self.assertEqual(replies[0][1].count('XACNHAN '),1)


if __name__=='__main__':
    unittest.main()
